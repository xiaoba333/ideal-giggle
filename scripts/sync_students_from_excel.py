# -*- coding: utf-8 -*-
"""一次性：删除 address/remark，并从学院花名册同步 2506 班学生电话与班级。"""

import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import openpyxl

from database.db_conn import execute_query, execute_update
from dao import student_dao

EXCEL = Path(r"C:\Users\LMY\Desktop\班级管理\计算机学院学生信息表-软件工程2506班.xlsx")


def drop_unused_columns():
    cols = execute_query(
        "SELECT COLUMN_NAME FROM information_schema.COLUMNS "
        "WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME='student' "
        "AND COLUMN_NAME IN ('address','remark')"
    )
    names = {c["COLUMN_NAME"] for c in cols}
    print("drop columns:", names)
    if "address" in names:
        execute_update("ALTER TABLE student DROP COLUMN address")
    if "remark" in names:
        execute_update("ALTER TABLE student DROP COLUMN remark")


def sync_from_excel():
    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    ws = wb.active
    updated = inserted = 0
    for r in range(2, ws.max_row + 1):
        cls = str(ws.cell(r, 2).value or "")
        if "2506" not in cls:
            continue
        no = str(ws.cell(r, 4).value or "").strip()
        name = str(ws.cell(r, 3).value or "").strip()
        gender = ws.cell(r, 6).value
        gender = str(gender).strip() if gender else None
        phone = ws.cell(r, 11).value
        phone = str(phone).strip() if phone else None
        class_name = cls.strip() or None
        if not no or not name:
            continue
        exist = student_dao.get_student_by_no(no)
        if exist:
            student_dao.update_student(exist["id"], no, name, gender, class_name, phone)
            updated += 1
        else:
            student_dao.insert_student(no, name, gender, class_name, phone)
            inserted += 1
    return updated, inserted


def main():
    drop_unused_columns()
    updated, inserted = sync_from_excel()
    rows = execute_query(
        "SELECT student_no, name, gender, class_name, phone "
        "FROM student ORDER BY student_no"
    )
    print(f"updated={updated} inserted={inserted} total={len(rows)}")
    for row in rows:
        print(row)
    schema = execute_query(
        "SELECT COLUMN_NAME FROM information_schema.COLUMNS "
        "WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME='student' "
        "ORDER BY ORDINAL_POSITION"
    )
    print("columns:", [c["COLUMN_NAME"] for c in schema])


if __name__ == "__main__":
    main()
