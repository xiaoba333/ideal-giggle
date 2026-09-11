# -*- coding: utf-8 -*-
"""
从学院花名册重建学生表：
- 去掉班级字段
- 导入姓名/学号/性别/手机
- id 从 1 连续重排
"""

import json
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import openpyxl

from database.db_conn import execute_query, execute_update
from dao import student_dao
from dao.classlog_dao import dumps_participants, loads_participants

EXCEL = Path(r"C:\Users\LMY\Desktop\班级管理\计算机学院学生信息表-软件工程2506班.xlsx")


def load_excel_students():
    """班级含 2506 的学生；学号去重（保留先出现的）。"""
    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    ws = wb.active
    rows = []
    seen = set()
    skipped = []
    for r in range(2, ws.max_row + 1):
        cls = str(ws.cell(r, 2).value or "")
        no = str(ws.cell(r, 4).value or "").strip()
        name = str(ws.cell(r, 3).value or "").strip()
        gender = ws.cell(r, 6).value
        gender = str(gender).strip() if gender else None
        phone = ws.cell(r, 11).value
        phone = str(phone).strip() if phone else None
        if not no or not name:
            continue
        if "2506" not in cls:
            continue
        if no in seen:
            skipped.append((no, name, cls))
            continue
        seen.add(no)
        rows.append({
            "student_no": no,
            "name": name,
            "gender": gender,
            "phone": phone,
        })
    # 若不足 27，再补入学号以 82092506 开头且尚未收录的行（如花名册误写班级）
    if len(rows) < 27:
        for r in range(2, ws.max_row + 1):
            no = str(ws.cell(r, 4).value or "").strip()
            name = str(ws.cell(r, 3).value or "").strip()
            cls = str(ws.cell(r, 2).value or "")
            gender = ws.cell(r, 6).value
            gender = str(gender).strip() if gender else None
            phone = ws.cell(r, 11).value
            phone = str(phone).strip() if phone else None
            if not no or not name or not no.startswith("82092506"):
                continue
            if no in seen:
                continue
            seen.add(no)
            rows.append({
                "student_no": no,
                "name": name,
                "gender": gender,
                "phone": phone,
            })
            if len(rows) >= 27:
                break
    # 仍不足且存在重复学号候选（如蒋书涵）：用缺失学号 8209250622 收录
    if len(rows) < 27:
        for r in range(2, ws.max_row + 1):
            no = str(ws.cell(r, 4).value or "").strip()
            name = str(ws.cell(r, 3).value or "").strip()
            cls = str(ws.cell(r, 2).value or "")
            gender = ws.cell(r, 6).value
            gender = str(gender).strip() if gender else None
            phone = ws.cell(r, 11).value
            phone = str(phone).strip() if phone else None
            if name != "蒋书涵":
                continue
            alt_no = "8209250622"
            if alt_no in seen:
                break
            rows.append({
                "student_no": alt_no,
                "name": name,
                "gender": gender,
                "phone": phone,
            })
            skipped.append((no, f"{name}->改用{alt_no}", cls))
            break
    rows.sort(key=lambda x: x["student_no"])
    return rows, skipped


def remap_log_participants(name_to_id):
    logs = execute_query("SELECT id, participants FROM class_log")
    for log in logs:
        people = loads_participants(log.get("participants"))
        if not people:
            continue
        new_people = []
        for p in people:
            name = (p.get("name") or "").strip()
            sid = name_to_id.get(name)
            item = {"name": name}
            if sid is not None:
                item["id"] = sid
            elif p.get("id") is not None:
                item["id"] = p["id"]
            if item.get("id") is not None or name:
                new_people.append(item)
        execute_update(
            "UPDATE class_log SET participants=%s WHERE id=%s",
            (dumps_participants(new_people), log["id"]),
        )


def main():
    student_dao.ensure_schema()
    students, skipped = load_excel_students()
    print(f"excel students={len(students)} skipped={skipped}")

    execute_update("DELETE FROM student")
    for i, stu in enumerate(students, start=1):
        student_dao.insert_student(
            stu["student_no"],
            stu["name"],
            stu.get("gender"),
            stu.get("phone"),
            student_id=i,
        )
    next_id = len(students) + 1
    execute_update(f"ALTER TABLE student AUTO_INCREMENT = {next_id}")

    rows = execute_query(
        "SELECT id, student_no, name, gender, phone FROM student ORDER BY id"
    )
    name_to_id = {r["name"]: r["id"] for r in rows}
    remap_log_participants(name_to_id)

    print(f"imported={len(rows)} ids=1..{len(rows)} next_auto={next_id}")
    for r in rows:
        print(r)
    cols = execute_query(
        "SELECT COLUMN_NAME FROM information_schema.COLUMNS "
        "WHERE TABLE_SCHEMA=DATABASE() AND TABLE_NAME='student' "
        "ORDER BY ORDINAL_POSITION"
    )
    print("columns:", [c["COLUMN_NAME"] for c in cols])


if __name__ == "__main__":
    main()
