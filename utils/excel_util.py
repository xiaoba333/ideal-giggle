# -*- coding: utf-8 -*-
"""openpyxl 学生数据 Excel 导入导出（含班级、初始德育分）。"""

from openpyxl import Workbook, load_workbook

STUDENT_HEADERS = [
    "学号",
    "姓名",
    "性别",
    "联系电话",
    "班级",
    "初始德育分",
]

HEADER_TO_FIELD = {
    "学号": "student_no",
    "姓名": "name",
    "性别": "gender",
    "联系电话": "phone",
    "联系方式": "phone",
    "班级": "class_name",
    "初始德育分": "moral_base",
    "德育初始分": "moral_base",
    "德育分": "moral_base",
}


def export_students_to_excel(students, file_path):
    """
    导出学生列表到 Excel。
    :param students: list[dict]
    :param file_path: 目标 .xlsx 路径
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "学生信息"
    ws.append(STUDENT_HEADERS)

    for row in students:
        ws.append([
            row.get("student_no") or "",
            row.get("name") or "",
            row.get("gender") or "",
            row.get("phone") or "",
            row.get("class_name") or "",
            row.get("moral_base") if row.get("moral_base") is not None else "",
        ])

    wb.save(file_path)


def import_students_from_excel(file_path):
    """
    从 Excel 导入学生数据。
    :return: list[dict]
    """
    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        wb.close()
        raise ValueError("Excel 文件为空")

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    if "学号" not in headers or "姓名" not in headers:
        wb.close()
        raise ValueError("Excel 必须包含「学号」「姓名」列")

    col_map = {}
    for idx, h in enumerate(headers):
        if h in HEADER_TO_FIELD:
            field = HEADER_TO_FIELD[h]
            if field == "phone" and "phone" in col_map and h == "联系方式":
                continue
            col_map[field] = idx

    result = []
    seen_nos = set()
    for row in rows_iter:
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue

        def cell(field):
            i = col_map.get(field)
            if i is None or i >= len(row):
                return None
            val = row[i]
            if val is None:
                return None
            if field == "moral_base":
                try:
                    return int(val)
                except (TypeError, ValueError):
                    text = str(val).strip()
                    if not text:
                        return None
                    return int(float(text))
            text = str(val).strip()
            return text if text else None

        student_no = cell("student_no")
        name = cell("name")
        if not student_no or not name:
            continue
        if student_no in seen_nos:
            continue
        seen_nos.add(student_no)

        item = {
            "student_no": student_no,
            "name": name,
            "gender": cell("gender"),
            "phone": cell("phone"),
            "class_name": cell("class_name"),
            "moral_base": None,
        }
        try:
            item["moral_base"] = cell("moral_base")
        except (TypeError, ValueError):
            item["moral_base"] = None
        result.append(item)

    wb.close()
    return result
